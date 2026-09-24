#!/usr/bin/env python3
"""Build aligned editions and a portable, source-based arXiv candidate package.
Only --format repairs narrowly specified Markdown punctuation; --check is read-only.
The tests validate structure and the toy model, not translation semantics or scholarship.
"""
from __future__ import annotations
import argparse
import hashlib
import json
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from toy_authority import run as toy_run

ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / '.build'
REPORTS = ROOT / 'build'
SUBMISSION = ROOT / 'submission'
READER = 'markdown+tex_math_dollars+raw_tex+fenced_divs+autolink_bare_uris'
SECTION = re.compile(r'<a id="section-(\d+)"></a>')
REFANCHOR = re.compile(r'<a id="ref-(\d+)"></a>')
CITE = re.compile(r'\[(\d+)\]\(#ref-(\d+)\)')
CJK = re.compile(r'[\u3400-\u4dbf\u4e00-\u9fff]')
CSS = '''<style>
html{background:#fff;color:#202124}body{max-width:820px;margin:3rem auto;padding:0 1.4rem;font:18px/1.75 Georgia,"Noto Serif CJK SC",serif}h1,h2,h3{line-height:1.3;color:#141414}h1{font-size:1.55rem;margin-top:2rem}header h1{font-size:2.15rem}header{text-align:center;margin-bottom:2.5rem}.subtitle{font-size:1.25rem}a{color:#174b75;overflow-wrap:anywhere}nav{font:15px/1.5 system-ui,sans-serif;border-bottom:1px solid #ddd;padding-bottom:1rem}table{border-collapse:collapse;width:100%;font-size:16px;margin:1.4rem 0}th,td{border-bottom:1px solid #ddd;padding:.55rem;text-align:left}math[display="block"]{overflow-x:auto;padding:1rem 0}code{font-size:.85em;overflow-wrap:anywhere}p{margin:0 0 1em}@media print{nav{display:none}body{max-width:none;font-size:11pt;margin:0}a{color:inherit}}
</style>'''


def command(args, *, text=None, cwd=ROOT):
    result = subprocess.run([str(a) for a in args], cwd=cwd, input=text,
                            text=True, capture_output=True)
    if result.returncode:
        raise RuntimeError('Command failed: ' + ' '.join(map(str, args)) + '\n' +
                           result.stdout[-8000:] + result.stderr[-8000:])
    return result.stdout


def pandoc(text, target, extra=()):
    return command(['pandoc', '-f', READER, '-t', target, '--wrap=none', *extra], text=text)


def parse(source):
    prefix = source.split('<a id="section-1"></a>', 1)[0]
    title = re.search(r'^# (.+)$', prefix, re.M).group(1)
    subtitle = re.search(r'^## (.+)$', prefix, re.M).group(1)
    date = [s.strip() for s in prefix.splitlines() if s.strip()][-1]
    parts = SECTION.split(source)[1:]
    sections = {int(parts[i]): parts[i + 1].strip() for i in range(0, len(parts), 2)}
    refparts = REFANCHOR.split(sections[15])[1:]
    references = {}
    for i in range(0, len(refparts), 2):
        n = int(refparts[i])
        references[n] = re.sub(r'^\*\*\[\d+\]\*\*\s*', '', refparts[i+1].strip())
    return {'title': title, 'subtitle': subtitle, 'date': date,
            'sections': sections, 'references': references}


def load_pair(format_sources=False):
    sources = {}
    repairs = {}
    for lang in ('en', 'zh'):
        path = ROOT / f'paper.{lang}.md'
        original = path.read_text(encoding='utf-8')
        normalized, count = re.subn(r'\]\(#ref-(\d+)\]', r'](#ref-\1)', original)
        if count and not format_sources:
            raise ValueError(f'{path.name}: malformed citation closing brackets; run --format')
        if normalized != original:
            path.write_text(normalized, encoding='utf-8')
        repairs[lang] = count
        sources[lang] = normalized
    parsed = {lang: parse(s) for lang, s in sources.items()}
    if CJK.search(sources['en']):
        raise ValueError('The English manuscript contains CJK characters.')
    for lang, source in sources.items():
        assert SECTION.findall(source) == [str(i) for i in range(1, 16)], lang
        assert REFANCHOR.findall(source) == [str(i) for i in range(1, 23)], lang
        assert set(parsed[lang]['references']) == set(range(1, 23)), lang
        assert all(a == b for a, b in CITE.findall(source)), lang
        assert set(int(a) for a, _ in CITE.findall(source)) == set(range(1, 23)), lang
        assert not re.search(r'\b(TODO|TBD|PLACEHOLDER)\b', source), lang
        assert not re.search(r'\]\(#ref-\d+[^)\d]', source), lang
    assert parsed['en']['references'] == parsed['zh']['references'], 'Reference text mismatch'
    details = []
    for n in range(1, 16):
        a, b = (parsed[lang]['sections'][n] for lang in ('en', 'zh'))
        assert CITE.findall(a) == CITE.findall(b), f'Section {n}: citation sequence mismatch'
        maths = [re.findall(r'\$\$\s*(.*?)\s*\$\$', s, re.S) for s in (a, b)]
        assert maths[0] == maths[1], f'Section {n}: equations differ'
        assert a.count('$$') % 2 == 0 and b.count('$$') % 2 == 0, n
        details.append({'section': n, 'citation_occurrences': len(CITE.findall(a)),
                        'display_equations': len(maths[0])})
    report = {'structural_alignment': 'passed', 'sections': details,
              'references': 22, 'english_cjk_characters': 0,
              'citation_punctuation_repairs': repairs,
              'translation_semantics': 'requires human review; not proven by structural tests'}
    return sources, parsed, report


def split_heading(section):
    heading, _, text = section.partition('\n')
    return heading.removeprefix('## ').strip(), text.strip()


def body_markdown(paper, *, tex=False):
    parts = []
    first = 2 if tex else 1
    for n in range(first, 15):
        heading, text = split_heading(paper['sections'][n])
        if tex:
            # URL's obeyspaces mode preserves commands and permits safe path wrapping.
            text = re.sub(r'`([^`\n]+)`', lambda m: r'\path{' + m.group(1) + '}', text)
            text = CITE.sub(lambda m: r'\cite{ref-' + m.group(1) + '}', text)
        else:
            text = CITE.sub(lambda m: '[[' + m.group(1) + ']](#ref-' + m.group(1) + ')', text)
        parts.append(f'# {heading} {{#section-{n}}}\n\n{text}\n')
    if tex:
        bib = ['\\begin{thebibliography}{99}', '\\small']
        for n, ref in paper['references'].items():
            bib.extend([f'\\bibitem{{ref-{n}}}', pandoc(ref, 'latex').strip()])
        bib.append('\\end{thebibliography}')
        parts.append('\n'.join(bib))
    else:
        heading, _ = split_heading(paper['sections'][15])
        parts.append(f'# {heading} {{#section-15}}')
        for n, ref in paper['references'].items():
            parts.append(f'::: {{#ref-{n}}}\n\n**[{n}]** {ref}\n\n:::')
    return '\n\n'.join(parts) + '\n'


def style_docx(path, lang):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document(path)
    # Pandoc and python-docx versions differ in localized/display style names.
    # Resolve using normalized names/IDs rather than built-in name translation.
    styles = {}
    for item in doc.styles:
        for key in (item.name, item.style_id):
            styles[re.sub(r"\s+", "", key).casefold()] = item
    def get_style(name):
        key = re.sub(r"\s+", "", name).casefold()
        if key not in styles:
            raise ValueError(f"Required paragraph style absent: {name}")
        return styles[key]
    for section in doc.sections:
        section.page_width, section.page_height = Inches(8.2677), Inches(11.6929)
        section.top_margin = section.bottom_margin = Inches(0.90)
        section.left_margin = section.right_margin = Inches(0.95)
        footer = section.footer.paragraphs[0]
        footer.alignment = 1
        field = OxmlElement('w:fldSimple')
        field.set(qn('w:instr'), 'PAGE')
        footer._p.append(field)
    for style_name in ('Normal', 'Body Text', 'First Paragraph', 'Title', 'Subtitle', 'Author', 'Date', 'Heading 1', 'Heading 2'):
        if re.sub(r'\s+', '', style_name).casefold() not in styles:
            continue
        style = get_style(style_name)
        style.font.name = 'Times New Roman'
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor(0, 0, 0)
        rpr = style.element.get_or_add_rPr()
        fonts = rpr.find(qn('w:rFonts'))
        if fonts is None:
            fonts = OxmlElement('w:rFonts')
            rpr.append(fonts)
        fonts.set(qn('w:eastAsia'), 'Noto Serif CJK SC')
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.space_after = Pt(6)
    get_style('Title').font.size = Pt(21)
    get_style('Subtitle').font.size = Pt(13)
    get_style('Heading 1').font.size = Pt(14)
    get_style('Heading 1').font.bold = True
    get_style('Heading 1').paragraph_format.space_before = Pt(14)
    get_style('Heading 1').paragraph_format.keep_with_next = True
    in_refs = False
    for paragraph in doc.paragraphs:
        if paragraph.text in ('References', '参考文献'):
            in_refs = True
        elif in_refs:
            for run in paragraph.runs:
                run.font.size = Pt(10)
    for table in doc.tables:
        table.autofit = True
        if table.rows:
            trpr = table.rows[0]._tr.get_or_add_trPr()
            repeat = OxmlElement('w:tblHeader')
            trpr.append(repeat)
            for cell in table.rows[0].cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
    doc.core_properties.author = 'Junliang Zhou'
    doc.save(path)


def compile_tex(tex_path, destination, engine, workdir):
    workdir.mkdir(parents=True, exist_ok=True)
    for _ in range(2):
        command([engine, '-interaction=nonstopmode', '-halt-on-error',
                 '-no-shell-escape', '-output-directory', workdir, tex_path])
    log = (workdir / (tex_path.stem + '.log')).read_text(encoding='utf-8', errors='replace')
    if re.search(r'undefined references|Citation .* undefined|Missing character:', log):
        raise ValueError(f'Unresolved reference or missing glyph in {tex_path.name}')
    shutil.copy2(workdir / (tex_path.stem + '.pdf'), destination)
    warnings = re.findall(r'Overfull \\[hv]box[^\n]*', log)
    if warnings:
        raise ValueError(f'Layout overflow in {tex_path.name}: {warnings}')
    return warnings


def package(papers, report):
    SUBMISSION.mkdir(exist_ok=True)
    abstract = split_heading(papers['en']['sections'][1])[1].split('**Keywords:')[0].strip()
    metadata = {'title': papers['en']['title'] + ': ' + papers['en']['subtitle'],
                'authors': ['Junliang Zhou'], 'abstract': abstract,
                'suggested_primary_category': 'cs.CY', 'cross_lists': [],
                'comments': 'Conceptual analysis with a reproducible finite-state illustration; Chinese companion in ancillary files.',
                'license': None, 'arxiv_id': None,
                'status': 'prepared_not_submitted', 'author_final_approval': 'pending',
                'arxiv_server_compile': 'not tested',
                'source_compiler': 'pdfLaTeX', 'main_file': 'main.tex'}
    (SUBMISSION / 'metadata.json').write_text(json.dumps(metadata, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    (SUBMISSION / 'abstract.txt').write_text(abstract + '\n', encoding='utf-8')
    explanation = ('English main.tex is self-contained and uses pdfLaTeX.\n'
                   'The Chinese PDF is a translation companion, not a second study.\n'
                   'toy_authority.py and authority-results.json reproduce the finite-state illustration.\n'
                   'No font files, credentials, or independent arXiv submission are included.\n')
    with zipfile.ZipFile(SUBMISSION / 'arxiv-upload.zip', 'w', zipfile.ZIP_DEFLATED) as archive:
        archive.write(ROOT / 'paper.en.tex', 'main.tex')
        archive.write(ROOT / 'paper.zh.pdf', 'anc/paper.zh.pdf')
        archive.write(ROOT / 'scripts/toy_authority.py', 'anc/toy_authority.py')
        archive.write(REPORTS / 'authority-results.json', 'anc/authority-results.json')
        archive.writestr('anc/README.txt', explanation)
    with tempfile.TemporaryDirectory() as directory:
        target = Path(directory)
        with zipfile.ZipFile(SUBMISSION / 'arxiv-upload.zip') as archive:
            archive.extractall(target)
        for _ in range(2):
            command(['pdflatex', '-interaction=nonstopmode', '-halt-on-error',
                     '-no-shell-escape', 'main.tex'], cwd=target)
        log = (target / 'main.log').read_text(errors='replace')
        assert 'undefined references' not in log
        assert 'Missing character:' not in log
    report['isolated_submission_source_compile'] = 'passed locally; arXiv server not tested'
    return metadata


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--check', action='store_true')
    parser.add_argument('--format', action='store_true')
    parser.add_argument('--tex-only', action='store_true')
    parser.add_argument('--formats-only', action='store_true')
    args = parser.parse_args()
    sources, papers, report = load_pair(args.format)
    results = toy_run()
    report['toy_coalition_checks'] = results['total_coalition_checks']
    if args.check:
        print(json.dumps(report, indent=2))
        return
    WORK.mkdir(exist_ok=True)
    REPORTS.mkdir(exist_ok=True)
    (REPORTS / 'authority-results.json').write_text(json.dumps(results, indent=2) + '\n', encoding='utf-8')
    report['latex_overfull_warnings'] = {}
    for lang, paper in papers.items():
        abstract = split_heading(paper['sections'][1])[1]
        label = '**Keywords:' if lang == 'en' else '**关键词：'
        abstract, _, keyword_tail = abstract.partition(label)
        metadata = {'title': paper['title'], 'subtitle': paper['subtitle'],
                    'author': ['Junliang Zhou'], 'date': paper['date'],
                    'lang': 'en-US' if lang == 'en' else 'zh-CN'}
        meta_path = WORK / f'meta.{lang}.json'
        meta_path.write_text(json.dumps(metadata, ensure_ascii=False), encoding='utf-8')
        body = body_markdown(paper)
        if not args.tex_only:
            other = 'zh' if lang == 'en' else 'en'
            navtext = 'Chinese edition' if lang == 'en' else '英文版'
            nav = WORK / f'nav.{lang}.html'
            nav.write_text(CSS + f'<nav><a href="paper.{other}.html">{navtext}</a> · <a href="paper.{lang}.pdf">PDF</a> · <a href="paper.{lang}.docx">Word</a></nav>', encoding='utf-8')
            html = pandoc(body, 'html5', ['--standalone', '--mathml', '--metadata-file', meta_path, '--include-before-body', nav])
            (ROOT / f'paper.{lang}.html').write_text(html, encoding='utf-8')
            command(['pandoc', '-f', READER, '-t', 'docx', '--standalone', '--metadata-file', meta_path,
                     '-o', ROOT / f'paper.{lang}.docx'], text=body)
            style_docx(ROOT / f'paper.{lang}.docx', lang)
        if not args.formats_only:
            metadata['abstract'] = abstract.strip()
            metadata['chinese'] = lang == 'zh'
            meta_path.write_text(json.dumps(metadata, ensure_ascii=False), encoding='utf-8')
            tex_body = body_markdown(paper, tex=True)
            if keyword_tail:
                keylabel = '**Keywords:**' if lang == 'en' else '**关键词：**'
                clean_keywords = keyword_tail.removeprefix('**').removeprefix('*').strip()
                tex_body = keylabel + ' ' + clean_keywords + '\n\n' + tex_body
            tex = pandoc(tex_body, 'latex', ['--standalone', '--template', ROOT / 'templates/paper.tex', '--metadata-file', meta_path])
            tex_path = ROOT / f'paper.{lang}.tex'
            tex_path.write_text(tex, encoding='utf-8')
            if not args.tex_only:
                engine = 'pdflatex' if lang == 'en' else 'xelatex'
                warnings = compile_tex(tex_path, ROOT / f'paper.{lang}.pdf', engine, WORK / lang)
                report['latex_overfull_warnings'][lang] = warnings
    if not args.tex_only and not args.formats_only:
        package(papers, report)
    files = [ROOT / f'paper.{lang}.{ext}' for lang in ('en', 'zh') for ext in ('md', 'tex', 'pdf', 'html', 'docx')]
    report['sha256'] = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in files if p.exists()}
    report['tool_versions'] = {tool: command([tool, '--version']).splitlines()[0] for tool in ('pandoc', 'pdflatex', 'xelatex')}
    (REPORTS / 'validation.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
